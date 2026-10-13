#!/usr/bin/env python3
"""Tally how often each person appears across one or more attendance spreadsheets.

Layout assumed for every input sheet:
  * a header row contains "学生所属班级" (class), an optional "学生所属年级"
    (grade), and one or more columns whose header contains "学生姓名"
    (the per-day attendance / sign-in columns);
  * a sheet may stack several classes, each preceded by its own header row;
    rows whose class cell is not a number are ignored, which drops the
    repeated header rows and any trailing teacher row automatically.

Every name cell in every per-day column is counted, and all inputs are
aggregated by (name, class) into a single output table.

Example:
    python3 tally_attendance.py -o out.xlsx week1.xlsx week2.xlsx

Two names inside the same class that differ by a single character (e.g.
尹湉默 vs 尹湉黙) are likely the same person typed two ways. They are never
merged: each keeps its own count, the two output rows are highlighted, and a
notice describes the pair so a human can confirm it.
"""

import argparse
import collections
import html
import os
import re
import sys
import zipfile

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    sys.exit("openpyxl is required (pip install openpyxl)")

CLASS_HEADER = "学生所属班级"
GRADE_HEADER = "学生所属年级"
NAME_HINT = "学生姓名"
OUT_HEADERS = ("学生姓名", "行政班", "需要补考勤次数")
VARIANT_FILL = "FFF2CC"          # highlight for suspected same-person rows


def text(v):
    return "" if v is None else str(v).strip()


def parse_class(v):
    s = text(v)
    if not s:
        return None
    try:
        f = float(s)
    except ValueError:
        return None
    return int(f) if f == int(f) else None


def day_key(header):
    """Normalize a name-column header to just the date, e.g. '9月15日'."""
    m = re.search(r"(\d{1,2})\s*月\s*(\d{1,2})\s*日", header)
    if m:
        return "%d月%d日" % (int(m.group(1)), int(m.group(2)))
    return header.replace(NAME_HINT, "").strip() or header


# ---------------------------------------------------------------- readers ----

def sheets_openpyxl(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return [(ws.title, [list(r) for r in ws.values]) for ws in wb.worksheets]


def _col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _parse_sheet_xml(blob, shared):
    grid = {}
    maxr = maxc = -1
    for m in re.finditer(r"<c\b([^>]*?)(?:/>|>(.*?)</c>)", blob, re.S):
        attrs = dict(re.findall(r'([\w:.-]+)="([^"]*)"', m.group(1)))
        mm = re.match(r"([A-Z]+)(\d+)", attrs.get("r", ""))
        if not mm:
            continue
        c = _col_index(mm.group(1))
        r = int(mm.group(2)) - 1
        inner = m.group(2) or ""
        vm = re.search(r"<v>(.*?)</v>", inner, re.S)
        if vm is None:
            ts = re.findall(r"<t\b[^>]*>(.*?)</t>", inner, re.S)
            val = html.unescape("".join(ts)) if ts else None
        elif attrs.get("t") == "s":
            val = shared[int(vm.group(1))]
        else:
            val = html.unescape(vm.group(1))
        grid[(r, c)] = val
        maxr = max(maxr, r)
        maxc = max(maxc, c)
    return [[grid.get((r, c)) for c in range(maxc + 1)] for r in range(maxr + 1)]


def sheets_xml(path):
    """Independent reader: pull cell values straight from the worksheet XML."""
    z = zipfile.ZipFile(path)
    names = z.namelist()
    shared = []
    if "xl/sharedStrings.xml" in names:
        ss = z.read("xl/sharedStrings.xml").decode("utf-8", "replace")
        for m in re.finditer(r"<si\b[^>]*/>|<si\b[^>]*>(.*?)</si>", ss, re.S):
            shared.append(html.unescape("".join(
                re.findall(r"<t\b[^>]*>(.*?)</t>", m.group(1) or "", re.S))))
    relmap = {}
    rel_part = "xl/_rels/workbook.xml.rels"
    if rel_part in names:
        for m in re.finditer(r"<Relationship\b([^>]*)/?>",
                             z.read(rel_part).decode("utf-8", "replace")):
            a = dict(re.findall(r'([\w:.-]+)="([^"]*)"', m.group(1)))
            if "Id" in a and "Target" in a:
                relmap[a["Id"]] = a["Target"]
    wbxml = z.read("xl/workbook.xml").decode("utf-8", "replace")
    out = []
    for m in re.finditer(r"<sheet\b([^>]*)/?>", wbxml):
        a = dict(re.findall(r'([\w:.-]+)="([^"]*)"', m.group(1)))
        target = relmap.get(a.get("r:id", ""), "").lstrip("/")
        part = target if target.startswith("xl/") else "xl/" + target
        if part in names:
            out.append((a.get("name", part),
                        _parse_sheet_xml(z.read(part).decode("utf-8", "replace"), shared)))
    return out


# --------------------------------------------------------------- tallying ----

def detect_layout(rows):
    for i, row in enumerate(rows):
        cells = [text(c) for c in row]
        class_col = next((j for j, c in enumerate(cells) if CLASS_HEADER in c), None)
        if class_col is None:
            continue
        name_cols = [j for j, c in enumerate(cells) if NAME_HINT in c and j < class_col]
        if not name_cols:
            continue
        grade_col = next((j for j, c in enumerate(cells) if GRADE_HEADER in c), None)
        return {"name_cols": name_cols, "class_col": class_col,
                "grade_col": grade_col, "data_start": i + 1}
    return None


def tally(all_sheets):
    counts = collections.Counter()
    order = collections.OrderedDict()          # class -> [name, ...] (first seen)
    grades = {}
    notes = []
    warnings = []
    seen_cell = {}                             # (day, name, class) -> first src
    dup_cells = []
    per_src = collections.OrderedDict()        # src -> {"people": set, "occ": int}
    for src, sheet_name, rows in all_sheets:
        label = "%s / %s" % (src, sheet_name)
        stat = per_src.setdefault(src, {"people": set(), "occ": 0, "sheets": 0})
        layout = detect_layout(rows)
        if not layout:
            notes.append("%s: no header row found, skipped" % label)
            continue
        stat["sheets"] += 1
        nc, cc, gc, start = (layout["name_cols"], layout["class_col"],
                             layout["grade_col"], layout["data_start"])
        header = rows[start - 1]
        days = [day_key(text(header[j])) for j in nc]
        # row-major pass: counts
        for row in rows[start:]:
            if cc >= len(row):
                continue
            cls = parse_class(row[cc])
            if cls is None:
                continue
            if gc is not None and gc < len(row) and text(row[gc]):
                grades.setdefault(cls, text(row[gc]))
            for idx, j in enumerate(nc):
                nm = text(row[j])
                if nm:
                    counts[(nm, cls)] += 1
                    stat["people"].add((nm, cls))
                    stat["occ"] += 1
                    if seen_cell.setdefault((days[idx], nm, cls), src) != src:
                        dup_cells.append((days[idx], nm, cls))
        # column-major pass: roster-like first-appearance order
        for j in nc:
            for row in rows[start:]:
                if j >= len(row) or cc >= len(row):
                    continue
                cls = parse_class(row[cc])
                if cls is None:
                    continue
                nm = text(row[j])
                if nm and nm not in order.setdefault(cls, []):
                    order[cls].append(nm)
    for src, st in per_src.items():
        notes.append("%s: %d people, %d occurrences" % (src, len(st["people"]), st["occ"]))
    dup_cells = sorted(set(dup_cells))
    if dup_cells:
        sample = ", ".join("%s %s" % (d, n) for d, n, _ in dup_cells[:8])
        warnings.append(
            "%d student-day cell(s) appear in more than one input (%s%s); they are "
            "counted twice, so drop the redundant input" %
            (len(dup_cells), sample, " ..." if len(dup_cells) > 8 else ""))
    return counts, order, grades, notes, warnings


def build_rows(counts, order, grades):
    out = []
    for cls in sorted(order):
        label = "%s%s班" % (grades.get(cls, ""), cls)
        for nm in order[cls]:
            out.append((nm, label, counts[(nm, cls)]))
    return out


def _one_char_apart(a, b):
    if len(a) != len(b):
        return False
    return sum(1 for x, y in zip(a, b) if x != y) == 1


def find_variants(rows):
    """Flag look-alike names within the same 行政班.

    Two names in one class that differ by exactly one character are a likely
    typo for the same person, but they can also be two different students.
    We never merge them: both keep their own counts and are only highlighted
    for a human to confirm. Names in different classes are never compared.
    """
    by_class = collections.OrderedDict()
    for nm, label, _ in rows:
        by_class.setdefault(label, []).append(nm)
    flagged = set()
    notices = []
    for label, names in by_class.items():
        pairs = [(names[i], names[j])
                 for i in range(len(names))
                 for j in range(i + 1, len(names))
                 if _one_char_apart(names[i], names[j])]
        if not pairs:
            continue
        for a, b in pairs:
            flagged.add((a, label))
            flagged.add((b, label))
        detail = "、".join("「%s」/「%s」" % p for p in pairs)
        notices.append(
            "%s 里 %s 名字只差一个字，可能是同一个人在不同表里的不同写法。"
            "这里按两个名字分别计数，这两行在结果表中已用颜色标出，请核对。" %
            (label, detail))
    return flagged, notices


# ---------------------------------------------------------------- writing ----

def write_book(rows, path, flagged=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "考勤次数统计"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(list(OUT_HEADERS))
    for c in range(1, 4):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="微软雅黑", size=11, bold=True)
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for i, (nm, label, cnt) in enumerate(rows, start=2):
        fill = PatternFill("solid", fgColor=VARIANT_FILL) if (nm, label) in flagged else None
        ws.cell(row=i, column=1, value=nm)
        ws.cell(row=i, column=2, value=label)
        ws.cell(row=i, column=3, value=cnt)
        ws.cell(row=i, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=i, column=2).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=i, column=3).alignment = Alignment(horizontal="center", vertical="center")
        for c in range(1, 4):
            cell = ws.cell(row=i, column=c)
            cell.font = Font(name="微软雅黑", size=11)
            cell.border = border
            if fill is not None:
                cell.fill = fill
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 18
    ws.freeze_panes = "A2"
    wb.save(path)


def read_back(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(ws.values)
    got = collections.Counter()
    for r in rows[1:]:
        if r and r[0]:
            got[(r[0], r[1])] = int(r[2])
    return got, list(rows[0])


# ------------------------------------------------------------------- main ----

def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+", help="one or more .xlsx attendance files")
    ap.add_argument("-o", "--output", help="output .xlsx path")
    ap.add_argument("--no-verify", action="store_true",
                    help="skip the independent re-count cross-check")
    args = ap.parse_args()

    out = args.output or os.path.join(
        os.path.dirname(os.path.abspath(args.inputs[0])), "考勤次数统计.xlsx")

    all_sheets = []
    for p in args.inputs:
        for title, rows in sheets_openpyxl(p):
            all_sheets.append((os.path.basename(p), title, rows))

    counts, order, grades, notes, warnings = tally(all_sheets)
    rows = build_rows(counts, order, grades)
    flagged, notices = find_variants(rows)
    for n in notes:
        print("  " + n)
    if not rows:
        sys.exit("no attendance rows found - check that the sheets have a '%s' "
                 "column and '%s' name columns" % (CLASS_HEADER, NAME_HINT))
    write_book(rows, out, flagged)

    total = sum(c for _, _, c in rows)
    print("people: %d   occurrences: %d" % (len(rows), total))
    for w in warnings:
        print("  WARNING: " + w)
    for n in notices:
        print("  提示：" + n)
    print("saved: %s" % out)

    if args.no_verify:
        return

    xml_sheets = []
    for p in args.inputs:
        for title, rows_ in sheets_xml(p):
            xml_sheets.append((os.path.basename(p), title, rows_))
    counts2, _, _, _, _ = tally(xml_sheets)
    if counts2 != counts:
        diff = {k: (counts.get(k), counts2.get(k)) for k in set(counts) | set(counts2)
                if counts.get(k) != counts2.get(k)}
        sys.exit("VERIFY FAILED (openpyxl vs raw XML): %s" % diff)

    got, header = read_back(out)
    want = collections.Counter({(nm, label): c for nm, label, c in rows})
    if got != want:
        diff = {k: (want.get(k), got.get(k)) for k in set(want) | set(got)
                if want.get(k) != got.get(k)}
        sys.exit("VERIFY FAILED (written file): %s" % diff)
    if tuple(header) != OUT_HEADERS:
        sys.exit("VERIFY FAILED: header %r" % (header,))
    print("verified: openpyxl == raw XML == written file")


if __name__ == "__main__":
    main()
